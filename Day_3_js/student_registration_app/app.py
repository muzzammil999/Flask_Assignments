import io
import os
from datetime import datetime

import pymysql
from flask import (Flask, Response, jsonify, render_template, request,
                   send_file, url_for)
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image
from werkzeug.security import generate_password_hash

app = Flask(__name__)

# ---------------------------------------------------------------
# 1. MySQL connection settings  -> CHANGE THE PASSWORD HERE
# ---------------------------------------------------------------
DB_CONFIG = {
    "host": os.getenv("DB_HOST", "localhost"),
    "user": os.getenv("DB_USER", "root"),
    "password": os.getenv("DB_PASSWORD", "Chinna@9999"),
    "database": os.getenv("DB_NAME", "student_db"),
    "cursorclass": pymysql.cursors.DictCursor,
    "charset": "utf8mb4",
}

ALLOWED_PHOTO_EXTENSIONS = {"jpg", "jpeg", "png"}
MAX_PHOTO_SIZE = 2 * 1024 * 1024  # 2 MB


def get_db():
    """Open a new MySQL connection."""
    return pymysql.connect(**DB_CONFIG)


def allowed_photo(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_PHOTO_EXTENSIONS


# ---------------------------------------------------------------
# 2. Pages
# ---------------------------------------------------------------
@app.route("/", methods=["GET"])
def register_page():
    return render_template("register.html", form={})


@app.route("/register", methods=["POST"])
def register():
    form = request.form

    full_name = form.get("full_name", "").strip()
    email = form.get("email", "").strip()
    mobile = form.get("mobile", "").strip()
    dob = form.get("dob", "").strip()
    gender = form.get("gender", "").strip()
    course = form.get("course", "").strip()
    qualification = form.get("qualification", "").strip()
    training_mode = form.get("training_mode", "").strip()
    skills = form.getlist("skills")
    experience = form.get("experience", "").strip()
    address = form.get("address", "").strip()
    website = form.get("website", "").strip()
    password = form.get("password", "")
    confirm_password = form.get("confirm_password", "")
    agree = form.get("agree")

    required = [full_name, email, mobile, dob, gender, course,
                qualification, training_mode, address, password, confirm_password]
    if not all(required):
        return jsonify({"error": "Please fill in all required fields."}), 400
    if not skills:
        return jsonify({"error": "Please select at least one skill."}), 400
    if not agree:
        return jsonify({"error": "You must agree to register."}), 400
    if password != confirm_password:
        return jsonify({"error": "Passwords do not match."}), 400

    try:
        experience_value = int(experience) if experience else 0
        if experience_value < 0:
            raise ValueError
    except ValueError:
        return jsonify({"error": "Years of experience must be a non-negative number."}), 400

    # -----------------------------------------------------------
    # 3. IMAGE -> bytes.  This is the whole "trick":
    #    photo.read() gives the raw bytes of the image, and MySQL
    #    stores those bytes in the MEDIUMBLOB column.
    # -----------------------------------------------------------
    photo_bytes = None
    photo_type = None
    photo = request.files.get("photo")
    if photo and photo.filename:
        if not allowed_photo(photo.filename):
            return jsonify({"error": "Photo must be a JPEG or PNG file."}), 400

        photo_bytes = photo.read()                      # <-- image as bytes
        if len(photo_bytes) > MAX_PHOTO_SIZE:
            return jsonify({"error": "Photo size must be under 2 MB."}), 400

        try:                                            # make sure it's a real image
            Image.open(io.BytesIO(photo_bytes)).verify()
        except Exception:
            return jsonify({"error": "The uploaded file is not a valid image."}), 400

        photo_type = "image/png" if photo.filename.lower().endswith("png") else "image/jpeg"

    sql = """
        INSERT INTO students
            (full_name, email, mobile, dob, gender, course, qualification,
             training_mode, skills, experience, address, website,
             photo, photo_type, password_hash)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
    """
    values = (full_name, email, mobile, dob, gender, course, qualification,
              training_mode, ", ".join(skills), experience_value, address,
              website or None, photo_bytes, photo_type,
              generate_password_hash(password))

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(sql, values)
        conn.commit()
    except pymysql.err.IntegrityError:
        return jsonify({"error": "This email is already registered."}), 400
    except pymysql.MySQLError as e:
        return jsonify({"error": f"Database error: {e}"}), 500
    finally:
        conn.close()

    return jsonify({"redirect_url": url_for("student_list")})


@app.route("/students", methods=["GET"])
def student_list():
    conn = get_db()
    try:
        with conn.cursor() as cur:
            # We do NOT select the photo bytes here (too heavy);
            # we only ask "does this student have a photo?"
            cur.execute("""
                SELECT id, full_name, email, mobile, dob, gender, course,
                       qualification, training_mode, skills, experience,
                       address, website, (photo IS NOT NULL) AS has_photo
                FROM students ORDER BY id DESC
            """)
            students = cur.fetchall()
    finally:
        conn.close()
    return render_template("students.html", students=students)


# ---------------------------------------------------------------
# 4. Show the image: read bytes from MySQL and send them to the browser
#    <img src="/photo/5"> calls this route.
# ---------------------------------------------------------------
@app.route("/photo/<int:student_id>")
def student_photo(student_id):
    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT photo, photo_type FROM students WHERE id=%s", (student_id,))
            row = cur.fetchone()
    finally:
        conn.close()

    if not row or not row["photo"]:
        return Response(status=404)
    return Response(row["photo"], mimetype=row["photo_type"] or "image/jpeg")


# ---------------------------------------------------------------
# 5. Export everything (with photos) to an Excel file
# ---------------------------------------------------------------
@app.route("/export")
def export_excel():
    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute("""
                SELECT id, full_name, email, mobile, dob, gender, course,
                       qualification, training_mode, skills, experience,
                       address, website, photo, created_at
                FROM students ORDER BY id
            """)
            rows = cur.fetchall()
    finally:
        conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = ["ID", "Full Name", "Email", "Mobile", "DOB", "Gender", "Course",
               "Qualification", "Training Mode", "Skills", "Experience (yrs)",
               "Address", "Website", "Registered On", "Photo"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [6, 22, 28, 15, 12, 12, 20, 15, 14, 24, 16, 35, 28, 20, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    for r, s in enumerate(rows, start=2):
        ws.append([s["id"], s["full_name"], s["email"], s["mobile"], s["dob"],
                   s["gender"], s["course"], s["qualification"], s["training_mode"],
                   s["skills"], s["experience"], s["address"], s["website"],
                   s["created_at"]])
        ws.row_dimensions[r].height = 75

        if s["photo"]:                                   # put the picture in the cell
            img = Image.open(io.BytesIO(s["photo"])).convert("RGB")
            img.thumbnail((90, 90))
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            buf.seek(0)
            ws.add_image(XLImage(buf), f"O{r}")

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    name = f"students_{datetime.now():%Y%m%d_%H%M}.xlsx"
    return send_file(out, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(debug=True)
